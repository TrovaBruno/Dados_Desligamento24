import os
import openpyxl
from openpyxl import load_workbook

def criar_ou_abrir_planilha(nome_arquivo):
    if os.path.exists(nome_arquivo):
        wb = load_workbook(nome_arquivo)
        planilha = wb.active
    else:
        wb = openpyxl.Workbook()
        planilha = wb.active
        planilha.title = "Dados de Desligamento"
        cabecalho = ["Nome", "Gestor", "Telefone", "Data do Desligamento", "Área"]
        planilha.append(cabecalho)
    return wb, planilha

def adicionar_dados(planilha, nome, gestor, telefone, data_desligamento, area):
    planilha.append([nome, gestor, telefone, data_desligamento, area])

def salvar_planilha(wb, nome_arquivo):
    wb.save(nome_arquivo)
    print(f"Os dados foram salvos no arquivo '{nome_arquivo}'.")

def main():
    # Especificar um diretório diferente para salvar o arquivo
    diretorio_salvar = "C:\\Users\\Usuario\\Documents"
    nome_arquivo = os.path.join(diretorio_salvar, "dados_desligamento.xlsx")

    wb, planilha = criar_ou_abrir_planilha(nome_arquivo)

    while True:
        # Solicitar dados do usuário
        nome = input("Digite o nome: ")
        gestor = input("Digite o gestor: ")
        telefone = input("Digite o telefone: ")
        data_desligamento = input("Digite a data do desligamento: ")  # Aceita qualquer valor
        area = input("Digite a área: ")

        # Adicionar dados à planilha
        adicionar_dados(planilha, nome, gestor, telefone, data_desligamento, area)

        # Perguntar ao usuário se deseja adicionar mais informações
        resposta = input("Deseja adicionar mais informações? (S/N): ").upper()
        if resposta != 'S':
            break

    # Salvar a planilha
    salvar_planilha(wb, nome_arquivo)

if __name__ == "__main__":
    main()
